"""Builds the licence-safe sample files in test/fixtures/.

Texts are public domain (Project Gutenberg #1342 Pride and Prejudice, #1228 On the Origin
of Species), with every Gutenberg header, footer and licence line removed. Everything else
is generated here. Needs: reportlab python-docx python-pptx openpyxl xlwt odfpy pillow.

    python3 tool/make_fixtures.py <pp.txt> <origin.txt>

v3 fixtures (formats, comics, DRM samples) need only the standard library and bsdtar:

    python3 tool/make_fixtures.py --v3

yellow-wallpaper.mobi and .azw3 are Project Gutenberg #1952's own Kindle files (public domain),
downloaded as-is: https://www.gutenberg.org/ebooks/1952.kindle.noimages and .kf8.images.
"""
import io
import os
import re
import sys
import zipfile


OUT = os.path.join(os.path.dirname(__file__), '..', 'test', 'fixtures')
os.makedirs(OUT, exist_ok=True)


def paragraphs(block):
    paras = []
    for p in re.split(r'\n\s*\n', block):
        p = ' '.join(l.strip() for l in p.strip().splitlines())
        if not p or p.startswith('[Illustration') or p.startswith('[_') or p.endswith(']]'):
            continue
        paras.append(p)
    return paras


def pp_chapters(path, count):
    text = open(path, encoding='utf-8').read()
    start = text.index('It is a truth universally acknowledged')
    text = text[start:]
    parts = re.split(r'\n\s*CHAPTER [IVXL]+\.\s*\n', text)
    chapters = []
    for i, part in enumerate(parts[:count]):
        paras = [re.sub(r'_(.+?)_', r'\1', p) for p in paragraphs(part)]
        chapters.append((f'Chapter {i + 1}', paras))
    return chapters


def origin_chapter(path):
    text = open(path, encoding='utf-8').read()
    start = text.index('CHAPTER III.\nSTRUGGLE FOR EXISTENCE.')
    end = text.index('CHAPTER IV.', start + 100)
    return paragraphs(text[start + len('CHAPTER III.\nSTRUGGLE FOR EXISTENCE.'):end])


def font(size):
    for f in ['/System/Library/Fonts/Supplemental/Georgia.ttf', '/System/Library/Fonts/Times.ttc']:
        if os.path.exists(f):
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def plate(w, h, label):
    im = Image.new('RGB', (w, h))
    d = ImageDraw.Draw(im)
    for y in range(h):
        t = y / h
        d.line([(0, y), (w, y)], fill=(int(200 - 90 * t), int(170 - 60 * t), int(110 - 30 * t)))
    d.ellipse([w * 0.3, h * 0.25, w * 0.7, h * 0.75], fill=(70, 90, 60))
    d.text((12, h - 30), label, fill=(250, 250, 245), font=font(18))
    return im


# ---------------------------------------------------------------- EPUB

def make_epub(chapters):
    cover = Image.new('RGB', (600, 900), (99, 57, 54))
    d = ImageDraw.Draw(cover)
    d.text((60, 300), 'Pride and', fill=(255, 232, 230), font=font(72))
    d.text((60, 390), 'Prejudice', fill=(255, 232, 230), font=font(72))
    d.line([(60, 500), (180, 500)], fill=(255, 232, 230), width=3)
    d.text((60, 530), 'JANE AUSTEN', fill=(255, 232, 230), font=font(32))
    buf = io.BytesIO()
    cover.save(buf, 'JPEG', quality=85)

    illo = io.BytesIO()
    plate(800, 500, 'He came down to see the place').save(illo, 'JPEG', quality=80)

    def xhtml(title, body):
        return ('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
                '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en">'
                f'<head><title>{title}</title><link rel="stylesheet" href="style.css"/></head><body>{body}</body></html>')

    files = {}
    manifest, spine, nav = [], [], []
    for i, (title, paras) in enumerate(chapters):
        body = f'<section epub:type="chapter"><h2>{title}</h2>'
        for j, p in enumerate(paras):
            p = p.replace('&', '&amp;').replace('<', '&lt;')
            if i == 0 and j == 0:
                p = p.replace('a single man in possession of a good fortune', '<em>a single man in possession of a good fortune</em>')
            body += f'<p>{p}</p>'
            if i == 0 and j == 8:
                body += '<figure><img src="images/illustration.jpg" alt="He came down to see the place"/><figcaption>He came down to see the place.</figcaption></figure>'
        if i == 1:
            body += '<blockquote><p>Mr. Bennet was so odd a mixture of quick parts, sarcastic humour, reserve, and caprice.</p></blockquote>'
            body += '<p>See also <a href="chapter1.xhtml">the opening chapter</a>.</p>'
        body += '</section>'
        name = f'chapter{i + 1}.xhtml'
        files[f'OEBPS/{name}'] = xhtml(title, body)
        manifest.append(f'<item id="c{i + 1}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i + 1}"/>')
        nav.append(f'<li><a href="{name}">{title}</a></li>')
    files['OEBPS/nav.xhtml'] = xhtml('Contents', '<nav epub:type="toc"><h1>Contents</h1><ol>' + ''.join(nav) + '</ol></nav>')
    files['OEBPS/style.css'] = 'body{font-family:serif;color:#333;background:#fff} h2{text-align:center;font-size:2em} p{text-indent:1.4em;margin:0}'
    opf = ('<?xml version="1.0" encoding="utf-8"?>\n<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:unfurl:fixture:pp</dc:identifier>'
           '<dc:title>Pride and Prejudice</dc:title><dc:creator>Jane Austen</dc:creator><dc:language>en</dc:language>'
           '<meta property="dcterms:modified">2026-10-02T00:00:00Z</meta><meta name="cover" content="cover"/></metadata>'
           '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
           '<item id="css" href="style.css" media-type="text/css"/>'
           '<item id="cover" href="images/cover.jpg" media-type="image/jpeg" properties="cover-image"/>'
           '<item id="illo" href="images/illustration.jpg" media-type="image/jpeg"/>' + ''.join(manifest) +
           '</manifest><spine>' + ''.join(spine) + '</spine></package>')
    files['OEBPS/content.opf'] = opf
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    path = os.path.join(OUT, 'pride-and-prejudice.epub')
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr(zipfile.ZipInfo('mimetype'), 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
        z.writestr('META-INF/container.xml', container, compress_type=zipfile.ZIP_DEFLATED)
        for k, v in files.items():
            z.writestr(k, v, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr('OEBPS/images/cover.jpg', buf.getvalue())
        z.writestr('OEBPS/images/illustration.jpg', illo.getvalue())


# ---------------------------------------------------------------- PDFs

def make_pdfs(origin_paras):
    from reportlab.lib.pagesizes import A5
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Image as RLImage
    from reportlab.pdfgen import canvas as rlcanvas

    body = ParagraphStyle('b', fontName='Times-Roman', fontSize=10.5, leading=14.5, alignment=TA_JUSTIFY, firstLineIndent=14)
    head = ParagraphStyle('h', fontName='Times-Bold', fontSize=15, leading=20, alignment=TA_CENTER, spaceAfter=10, spaceBefore=6)
    sub = ParagraphStyle('s', fontName='Times-Bold', fontSize=12, leading=16, alignment=TA_CENTER, spaceAfter=8, spaceBefore=14)
    cap = ParagraphStyle('c', fontName='Times-Italic', fontSize=8.5, leading=11, alignment=TA_CENTER, spaceAfter=8)
    plate_img = io.BytesIO()
    plate(900, 560, 'Plate 3').save(plate_img, 'JPEG', quality=80)

    class Doc(BaseDocTemplate):
        def afterFlowable(self, f):
            if isinstance(f, Paragraph) and f.style.name in ('h', 's'):
                key = f'k{self.seq.nextf("outline")}'
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(f.getPlainText(), key, level=0 if f.style.name == 'h' else 1)

    def on_page(c, doc):
        c.saveState()
        c.setFont('Times-Roman', 7.5)
        if doc.page > 1:
            c.drawString(42, A5[1] - 30, 'STRUGGLE FOR EXISTENCE')
            c.drawRightString(A5[0] - 42, A5[1] - 30, 'CHAP. III')
        c.drawCentredString(A5[0] / 2, 24, str(60 + doc.page))
        c.restoreState()

    def build(paras, path, encrypt=None):
        doc = Doc(path, pagesize=A5, leftMargin=42, rightMargin=42, topMargin=48, bottomMargin=44,
                  title='On the Origin of Species', author='Charles Darwin', encrypt=encrypt)
        doc.addPageTemplates([PageTemplate(id='p', frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height)], onPage=on_page)])
        story = [Paragraph('CHAPTER III.', head), Paragraph('STRUGGLE FOR EXISTENCE.', head)]
        for i, p in enumerate(paras):
            if i == 4:
                story.append(Paragraph('Geometrical Ratio of Increase.', sub))
            if i == 12:
                story.append(Paragraph('Nature of the Checks to Increase.', sub))
            story.append(Paragraph(p.replace('&', '&amp;'), body))
            if i == 6:
                plate_img.seek(0)
                story += [Spacer(1, 6), RLImage(plate_img, width=doc.width, height=doc.width * 0.62),
                          Paragraph('Plate 3. Geospiza, from the Galápagos Archipelago.', cap)]
        doc.build(story)
        return doc.page

    n = len(origin_paras)
    lo, hi = 1, n
    target = os.path.join(OUT, 'origin-of-species.pdf')
    while lo < hi:  # largest paragraph count that still fits in 10 pages
        mid = (lo + hi + 1) // 2
        if build(origin_paras[:mid], target) <= 10:
            lo = mid
        else:
            hi = mid - 1
    pages = build(origin_paras[:lo], target)
    print('origin pages', pages, 'paragraphs', lo)
    build(origin_paras[:lo], os.path.join(OUT, 'locked-statement.pdf'), encrypt='unfurl')

    # Two columns, four pages: for the extraction heuristics.
    from reportlab.lib.pagesizes import A4
    path = os.path.join(OUT, 'two-column.pdf')
    doc = BaseDocTemplate(path, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=60, bottomMargin=60, title='A survey of reflow', author='Kim and Osei')
    w = (doc.width - 24) / 2

    def head2(c, d):
        c.saveState()
        c.setFont('Times-Roman', 8)
        c.drawString(50, A4[1] - 40, 'Journal of Reading Interfaces · Vol. 12')
        c.drawCentredString(A4[0] / 2, 34, str(d.page))
        c.restoreState()

    doc.addPageTemplates([PageTemplate(id='two', frames=[Frame(doc.leftMargin, doc.bottomMargin, w, doc.height), Frame(doc.leftMargin + w + 24, doc.bottomMargin, w, doc.height)], onPage=head2)])
    hy = ParagraphStyle('hy', parent=body, fontSize=10, leading=13.5)
    story = [Paragraph('1. Introduction', sub)]
    for i, p in enumerate(origin_paras[:22]):
        story.append(Paragraph(p.replace('&', '&amp;'), hy))
        if i == 10:
            story.append(Paragraph('2. Method', sub))
    doc.build(story)

    # Scanned: two pages that are only images.
    path = os.path.join(OUT, 'field-notes-scanned.pdf')
    c = rlcanvas.Canvas(path, pagesize=A5)
    c.setTitle('Field Notes, Spring 2026')
    for k in range(2):
        im = Image.new('RGB', (840, 1190), (246, 242, 230))
        d = ImageDraw.Draw(im)
        for y in range(120, 1100, 44):
            d.line([(80, y), (780, y)], fill=(200, 205, 215))
            d.text((90, y - 30), 'observed finches near the shore, beak sizes vary' if (y // 44) % 3 else 'weather: dry, wind from the east', fill=(40, 50, 90), font=font(22))
        b = io.BytesIO()
        im.save(b, 'JPEG', quality=70)
        b.seek(0)
        from reportlab.lib.utils import ImageReader
        c.drawImage(ImageReader(b), 0, 0, width=A5[0], height=A5[1])
        c.showPage()
    c.save()


# ---------------------------------------------------------------- Office

def make_docx():
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.core_properties.title = 'Thesis draft v4'
    d.core_properties.author = 'A. Student'
    d.add_heading('Reading on small screens', 0)
    d.add_heading('1. Introduction', 1)
    d.add_paragraph('This draft looks at how people read long documents on phones, and what a reader can do to make fixed layouts comfortable on a small screen.')
    d.add_heading('2. Related work', 1)
    p = d.add_paragraph('Reading on small screens has been studied since the first dedicated e-readers. Most work compares ')
    p.add_run('paged').bold = True
    p.add_run(' and ')
    p.add_run('scrolling').italic = True
    p.add_run(' layouts for comprehension, with mixed results that depend on text length and task.')
    d.add_heading('2.1 Reflow of fixed layouts', 2)
    d.add_paragraph('Converting a fixed page into reflowable text requires recovering reading order from positioned glyphs. Table 1 summarises the approaches surveyed.')
    t = d.add_table(rows=4, cols=3)
    t.style = 'Table Grid'
    for r, row in enumerate([['Method', 'Order', 'Tables'], ['Glyph clustering', 'Good', 'Poor'], ['Tagged PDF', 'Exact', 'Good'], ['Hybrid', 'Good', 'Fair']]):
        for cidx, v in enumerate(row):
            t.cell(r, cidx).text = v
    cp = d.add_paragraph('Table 1. Reflow approaches.')
    cp.runs[0].italic = True
    cp.runs[0].font.size = Pt(9)
    d.add_paragraph('Where tags are absent, a hybrid approach falls back to geometry, which handles single columns well but struggles with floats and footnotes.')
    d.add_heading('3. Findings', 1)
    for s in ['Readers preferred their own font size over the original layout.', 'Sepia was the most used theme after dark.', 'Read aloud was used mostly while commuting.']:
        d.add_paragraph(s, style='List Bullet')
    for i in range(6):
        d.add_paragraph('A longer discussion paragraph follows so that the document runs onto a second page. ' * 4)
    d.save(os.path.join(OUT, 'Thesis draft v4.docx'))


def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.chart.data import CategoryChartData
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slides = [
        ('Q3 review', None, 'Title slide.'),
        ('What changed this quarter', ['Reader mode for PDF', 'Four reading themes'], 'Open with the reflow launch.'),
        ('Reader retention, Q3', ['30-day retention up 6 points', 'Reflow used in 41% of PDF sessions', 'Read aloud mostly on commutes'], 'Chart is 30-day retention by month.'),
        ('Next quarter', ['Tablet layouts', 'Faster large-file opening'], 'Keep this short.'),
        ('Thank you', ['Questions?'], ''),
    ]
    for i, (title, bullets, note) in enumerate(slides):
        layout = prs.slide_layouts[0 if bullets is None else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        if bullets is None:
            s.placeholders[1].text = 'Unfurl team · October 2026'
        else:
            body = s.placeholders[1].text_frame
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
        if i == 2:
            img = io.BytesIO()
            plate(600, 400, 'retention').save(img, 'PNG')
            img.seek(0)
            s.shapes.add_picture(img, Inches(8.5), Inches(2), width=Inches(4))
        if note:
            s.notes_slide.notes_text_frame.text = note
    prs.save(os.path.join(OUT, 'Q3 review.pptx'))


def make_sheets():
    import openpyxl
    rows = [['Month', 'Rent', 'Groceries', 'Utilities'], ['Mar', 1240, 412.30, 138.00], ['Apr', 1240, 398.75, 121.40],
            ['May', 1240, 441.10, 96.20], ['Jun', 1240, 405.60, 88.00], ['Jul', 1240, 462.95, 84.50], ['Aug', 1240, 430.00, 360.00],
            ['Sep', 1240, 419.80, 102.30], ['Oct', 1240, 388.15, 131.70]]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '2026'
    for r in rows:
        ws.append(r)
    ws.append(['Total', '=SUM(B2:B9)', '=SUM(C2:C9)', '=SUM(D2:D9)'])
    from openpyxl.comments import Comment
    ws['D7'].comment = Comment('Boiler service in August, paid half from savings.', 'me')
    ws2 = wb.create_sheet('Savings')
    for i in range(1, 400):
        ws2.append([f'Item {i}'] + [round(i * k * 1.37, 2) for k in range(1, 20)])
    wb.create_sheet('Notes').append(['Remember to check the gas bill'])
    wb.save(os.path.join(OUT, 'Household budget.xlsx'))

    import xlwt
    wb = xlwt.Workbook()
    s = wb.add_sheet('2025')
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            s.write(r, c, v)
    wb.save(os.path.join(OUT, 'Budget 2025.xls'))

    with open(os.path.join(OUT, 'survey-results.csv'), 'w') as f:
        f.write('respondent,theme,font,minutes per day,comment\n')
        themes = ['Light', 'Sepia', 'Dark', 'AMOLED']
        fonts = ['Literata', 'Instrument Sans', 'Atkinson', 'OpenDyslexic']
        for i in range(1, 201):
            f.write(f'{i},{themes[i % 4]},{fonts[(i * 7) % 4]},{(i * 13) % 90 + 5},"likes the {themes[i % 4].lower()} page, reads, mostly"\n')

    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    t = Table(name='Plan')
    for row in [['Week', 'Book', 'Pages'], ['1', 'Persuasion', '80'], ['2', 'Emma', '120'], ['3', 'Moby-Dick', '150']]:
        tr = TableRow()
        for v in row:
            num = v.isdigit()
            tc = TableCell(valuetype='float' if num else 'string', value=v if num else None)
            tc.addElement(P(text=v))
            tr.addElement(tc)
        t.addElement(tr)
    doc.spreadsheet.addElement(t)
    doc.save(os.path.join(OUT, 'reading-plan.ods'))


def make_text():
    open(os.path.join(OUT, 'Reading list.md'), 'w').write(
        '# Reading list\n\nBooks for the winter, roughly in order. Ask Sam about the [Beagle edition](https://example.org/beagle) with plates.\n\n'
        '## Novels\n\n- [x] Persuasion\n- [ ] Northanger Abbey\n- [ ] Bartleby, the Scrivener\n\n'
        '## Notes\n\n> Read Origin alongside the Beagle journal; the field notes come first.\n\n'
        '```\nch. 3, 4, 14 first\n```\n\n1. Darwin, *On the Origin of Species*\n2. Darwin, **The Voyage of the Beagle**\n\n---\n\n'
        '| Book | Pages |\n|---|---|\n| Persuasion | 249 |\n| Emma | 474 |\n')
    open(os.path.join(OUT, 'packing-list.txt'), 'w').write(
        'Packing, Lisbon\n\npassport and cards\ncharger, adaptor (type F)\ntwo books, or just the phone\nrain jacket\nwalking shoes\nsunglasses\nnotebook\nthe good pen\n')
    open(os.path.join(OUT, 'config-sample.txt'), 'w').write(
        'name     | size   | modified\n---------+--------+---------\nreader.py|  4 KB  | Mon\nlayout.c | 12 KB  | Tue\n'
        'def unfurl(page):\n    return [line.strip() for line in page]\n')


def make_images():
    im = Image.new('RGB', (900, 1200), (250, 250, 247))
    d = ImageDraw.Draw(im)
    d.text((80, 80), 'CAFÉ DO RIO', fill=(20, 20, 20), font=font(48))
    for i, (item, price) in enumerate([('Bica', '0.90'), ('Pastel de nata', '1.40'), ('Torrada', '2.10'), ('Sumo de laranja', '3.20')]):
        d.text((80, 220 + i * 70), item, fill=(30, 30, 30), font=font(34))
        d.text((700, 220 + i * 70), price, fill=(30, 30, 30), font=font(34))
    d.text((80, 600), 'TOTAL   7.60 EUR', fill=(0, 0, 0), font=font(40))
    im.save(os.path.join(OUT, 'Receipt_0412.jpg'), quality=85)
    plate(1280, 800, 'figure 1').save(os.path.join(OUT, 'figure-1.png'))
    plate(1000, 700, 'harbour').save(os.path.join(OUT, 'harbour.webp'))
    with zipfile.ZipFile(os.path.join(OUT, 'dataset.zip'), 'w') as z:
        z.writestr('data.txt', 'not a format Unfurl reads')


def v3():
    import base64
    import binascii
    import struct
    import subprocess
    import tempfile
    import zlib

    def png(w, h, hue, k=0, label=0):
        """A comic page: paper with four flat panels (layout k), no font needed."""
        layouts = [[(5, 4, 90, 36), (5, 43, 43, 27), (52, 43, 43, 27), (5, 73, 90, 23)],
                   [(5, 4, 43, 44), (52, 4, 43, 20), (52, 27, 43, 21), (5, 51, 90, 45)],
                   [(5, 4, 90, 22), (5, 29, 90, 32), (5, 64, 43, 32), (52, 64, 43, 32)]]
        import colorsys
        rows = [bytearray([247, 245, 240] * w) for _ in range(h)]
        for i, (x, y, pw, ph) in enumerate(layouts[k % 3]):
            r, g, b = (int(c * 255) for c in colorsys.hls_to_rgb(((hue + i * 40) % 360) / 360, 0.75 - i * 0.07, 0.35))
            x0, y0, x1, y1 = w * x // 100, h * y // 100, w * (x + pw) // 100, h * (y + ph) // 100
            for yy in range(y0, y1):
                row = rows[yy]
                for xx in range(x0, x1):
                    edge = yy in (y0, y1 - 1) or xx in (x0, x1 - 1)
                    row[xx * 3:xx * 3 + 3] = bytes((40, 40, 40) if edge else (r, g, b))
        # The page number as a row of ticks along the bottom, so pages differ.
        for t in range(label):
            for yy in range(h - 14, h - 6):
                rows[yy][(8 + t * 6) * 3:(8 + t * 6 + 3) * 3] = bytes([30, 30, 30] * 3)
        raw = b''.join(b'\0' + bytes(r) for r in rows)
        def chunk(t, d):
            return struct.pack('>I', len(d)) + t + d + struct.pack('>I', binascii.crc32(t + d) & 0xffffffff)
        return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
                + chunk(b'IDAT', zlib.compress(raw, 9)) + chunk(b'IEND', b''))

    def info(series, number=None, volume=None, writer=None, rtl=False, pages=0):
        f = [f'<Series>{series}</Series>']
        if number: f.append(f'<Number>{number}</Number>')
        if volume: f.append(f'<Volume>{volume}</Volume>')
        if writer: f.append(f'<Writer>{writer}</Writer>')
        f.append(f'<PageCount>{pages}</PageCount>')
        if rtl: f.append('<Manga>YesAndRightToLeft</Manga>')
        return ('<?xml version="1.0"?>\n<ComicInfo xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
                + ''.join(f) + '</ComicInfo>').encode()

    def pages(n, hue, names):
        return [(names(i), png(240, 360, hue + i * 25, i, i + 1)) for i in range(n)]

    # Unpadded names on purpose: natural sort must put p2 before p10.
    nemo = pages(12, 25, lambda i: f'nemo/p{i + 1}.png')
    with zipfile.ZipFile(os.path.join(OUT, 'little-nemo-14.cbz'), 'w', zipfile.ZIP_STORED) as z:
        z.writestr('ComicInfo.xml', info('Little Nemo in Slumberland', number='14', writer='Winsor McCay', pages=12))
        for name, data in reversed(nemo):
            z.writestr(name, data)
        z.writestr('__MACOSX/nemo/._p1.png', b'junk')

    # A .cbr that is really a zip: the magic bytes decide.
    with zipfile.ZipFile(os.path.join(OUT, 'mislabelled-zip.cbr'), 'w', zipfile.ZIP_STORED) as z:
        for name, data in pages(3, 120, lambda i: f'{i + 1:02}.png'):
            z.writestr(name, data)

    # A download cut short: the central directory is gone, the first pages survive.
    whole = io.BytesIO()
    with zipfile.ZipFile(whole, 'w', zipfile.ZIP_STORED) as z:
        for name, data in pages(6, 60, lambda i: f'page{i + 1}.png'):
            z.writestr(name, data)
    data = whole.getvalue()
    open(os.path.join(OUT, 'truncated.cbz'), 'wb').write(data[:len(data) * 2 // 3])

    # CBR: RAR 4 with stored entries, written by hand (no RAR tool is free to use).
    def rar4(entries):
        def block(htype, flags, body):
            head = struct.pack('<BHH', htype, flags, 7 + len(body)) + body
            return struct.pack('<H', binascii.crc32(head) & 0xffff) + head
        out = b'Rar!\x1a\x07\x00' + block(0x73, 0, b'\0' * 6)
        for name, data in entries:
            n = name.encode()
            body = struct.pack('<IIBIIBBHI', len(data), len(data), 0, binascii.crc32(data) & 0xffffffff,
                               0x5A2B0000, 29, 0x30, len(n), 0x20) + n
            out += block(0x74, 0x8000, body) + data
        return out + bytes.fromhex('c43d7b00400700')
    krazy = pages(5, 70, lambda i: f'Krazy Kat 1922/{i + 1:03}.png')
    open(os.path.join(OUT, 'krazy-kat-1922.cbr'), 'wb').write(
        rar4([('ComicInfo.xml', info('Krazy Kat', volume='1922', writer='George Herriman', pages=5))] + krazy))
    # RAR 5: only the signature and a header; must report "unsupported".
    open(os.path.join(OUT, 'rar5-sample.cbr'), 'wb').write(b'Rar!\x1a\x07\x01\x00' + b'\0' * 64)

    # CB7 and CBT through bsdtar (libarchive).
    with tempfile.TemporaryDirectory() as tmp:
        for name, data in pages(8, 200, lambda i: f'{i + 1}.png'):
            open(os.path.join(tmp, name), 'wb').write(data)
        open(os.path.join(tmp, 'ComicInfo.xml'), 'wb').write(
            info('Hokusai Manga', volume='2', writer='Katsushika Hokusai', rtl=True, pages=8))
        names = sorted(os.listdir(tmp))
        subprocess.run(['bsdtar', '--format', '7zip', '-cf', os.path.abspath(os.path.join(OUT, 'hokusai-manga-2.cb7'))] + names, cwd=tmp, check=True)
    with tempfile.TemporaryDirectory() as tmp:
        for name, data in pages(4, 300, lambda i: f'scan_{i + 1:04}.png'):
            open(os.path.join(tmp, name), 'wb').write(data)
        subprocess.run(['bsdtar', '--format', 'ustar', '-cf', os.path.abspath(os.path.join(OUT, 'scan_0412.cbt'))] + sorted(os.listdir(tmp)), cwd=tmp, check=True)

    # FB2 in Windows-1251 (Pushkin, 1833, public domain), a cover and a footnote.
    cover = base64.b64encode(png(120, 180, 30, 1)).decode()
    onegin = f'''<?xml version="1.0" encoding="windows-1251"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
<description><title-info><author><first-name>Александр</first-name><last-name>Пушкин</last-name></author>
<book-title>Евгений Онегин</book-title><coverpage><image l:href="#cover.png"/></coverpage><lang>ru</lang></title-info></description>
<body><title><p>Евгений Онегин</p></title>
<epigraph><p>Petri de vanite il avait encore plus de cette espece d'orgueil...</p><text-author>Tire d'une lettre particuliere</text-author></epigraph>
<section id="ch1"><title><p>Глава первая</p></title>
<poem><stanza><v>Мой дядя самых честных правил,</v><v>Когда не в шутку занемог,</v><v>Он уважать себя заставил</v><v>И лучше выдумать не мог.<a l:href="#n1" type="note">1</a></v></stanza></poem>
<p>Так думал молодой <emphasis>повеса</emphasis>, летя в пыли на почтовых.</p></section>
<section id="ch2"><title><p>Глава вторая</p></title><p>Деревня, где скучал Евгений, была прелестный уголок.</p><empty-line/><p><strong>Конец.</strong></p></section>
</body>
<body name="notes"><section id="n1"><title><p>1</p></title><p>Примечание автора.</p></section></body>
<binary id="cover.png" content-type="image/png">{cover}</binary>
</FictionBook>'''
    open(os.path.join(OUT, 'onegin.fb2'), 'wb').write(onegin.encode('cp1251', errors='replace'))

    walden_fb2 = '''<?xml version="1.0" encoding="utf-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
<description><title-info><author><first-name>Henry</first-name><middle-name>David</middle-name><last-name>Thoreau</last-name></author>
<book-title>Walden</book-title><lang>en</lang></title-info></description>
<body><section><title><p>Where I Lived, and What I Lived For</p></title>
<p>I went to the woods because I wished to live deliberately, to front only the essential facts of life, and see if I could not learn what it had to teach, and not, when I came to die, discover that I had not lived.</p>
<p>Our life is frittered away by detail. Simplify, simplify.</p></section></body></FictionBook>'''
    with zipfile.ZipFile(os.path.join(OUT, 'walden.fbz'), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('walden.fb2', walden_fb2)

    # Single-file HTML: unclosed paragraphs, a script, an inline image.
    dot = base64.b64encode(png(60, 40, 90)).decode()
    open(os.path.join(OUT, 'walden.html'), 'w').write(f'''<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Walden</title><meta name="author" content="Henry David Thoreau">
<style>p {{ color: red }}</style><script>document.write("tracking")</script></head>
<body><h1>Walden</h1><h2 id="economy">Economy</h2>
<p>When I wrote the following pages, or rather the bulk of them, I lived alone, in the woods, a mile from any neighbor.
<p>I should not talk so much about myself if there were anybody else whom I knew as well.<br>
<img src="data:image/png;base64,{dot}" alt="A pond">
<h2 id="where">Where I Lived</h2><p>Simplify, simplify. <a href="#economy">Back to Economy</a>
<ul><li>Shelter<li>Food<li>Fuel</ul>
</body></html>''')

    # DRM and variant samples: real-looking flags, no real protected content.
    src = os.path.join(OUT, 'pride-and-prejudice.epub')
    with zipfile.ZipFile(src) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    def epub_with(name, extra):
        with zipfile.ZipFile(os.path.join(OUT, name), 'w') as z:
            for i, d in items:
                z.writestr(i, d)
            for n, d in extra:
                z.writestr(n, d)
    enc = lambda alg, uri: f'''<?xml version="1.0"?><encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" xmlns:enc="http://www.w3.org/2001/04/xmlenc#"><enc:EncryptedData><enc:EncryptionMethod Algorithm="{alg}"/><enc:CipherData><enc:CipherReference URI="{uri}"/></enc:CipherData></enc:EncryptedData></encryption>'''
    epub_with('drm-sample.epub', [('META-INF/encryption.xml', enc('http://www.w3.org/2001/04/xmlenc#aes128-cbc', 'OEBPS/ch1.xhtml'))])
    epub_with('font-obfuscated.epub', [('META-INF/encryption.xml', enc('http://www.idpf.org/2008/embedding', 'OEBPS/fonts/a.otf'))])
    mobi = bytearray(open(os.path.join(OUT, 'yellow-wallpaper.mobi'), 'rb').read())
    r0 = struct.unpack('>I', mobi[78:82])[0]
    mobi[r0 + 12:r0 + 14] = struct.pack('>H', 2)  # Mobipocket DRM
    open(os.path.join(OUT, 'drm-sample.azw'), 'wb').write(bytes(mobi))
    open(os.path.join(OUT, 'kfx-sample.azw'), 'wb').write(b'\xeaDRMION\xee' + b'\0' * 248)
    open(os.path.join(OUT, 'topaz-sample.azw'), 'wb').write(b'TPZ0' + b'\0' * 252)


if __name__ == '__main__':
    if sys.argv[1:] == ['--v3']:
        v3()
        sys.exit(0)
    from PIL import Image, ImageDraw, ImageFont  # noqa: F401 (v1 fixtures only)
    pp, origin = sys.argv[1], sys.argv[2]
    chapters = pp_chapters(pp, 3)
    chapters[2] = (chapters[2][0], chapters[2][1][:8])  # about ten pages in all
    make_epub(chapters)
    make_pdfs(origin_chapter(origin))
    make_docx()
    make_pptx()
    make_sheets()
    make_text()
    make_images()
    print(sorted(os.listdir(OUT)))
